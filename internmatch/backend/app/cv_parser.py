"""Turn an uploaded CV (PDF / DOCX / TXT) into plain text."""
import io
import re
from pathlib import Path

import pdfplumber
from docx import Document

MAX_CHARS = 15_000
SUPPORTED_EXTENSIONS = (".pdf", ".docx", ".txt")


def extract_text(file_bytes: bytes, filename: str) -> str:
    """Return the CV's text, whitespace-collapsed and capped at MAX_CHARS.

    Raises ValueError with a user-facing message for unsupported, unreadable or empty files.
    """
    ext = Path(filename or "").suffix.lower()
    if ext not in SUPPORTED_EXTENSIONS:
        raise ValueError(f"Unsupported file type '{ext or 'unknown'}'. Please upload a PDF, DOCX or TXT file.")

    try:
        if ext == ".pdf":
            text = _pdf_text(file_bytes)
        elif ext == ".docx":
            text = _docx_text(file_bytes)
        else:
            text = _txt_text(file_bytes)
    except Exception as exc:  # corrupt / mislabelled files raise all sorts of parser errors
        raise ValueError(f"Couldn't read this {ext.lstrip('.').upper()} file. Is it corrupted?") from exc

    text = re.sub(r"\s+", " ", text).strip()
    if not text:
        raise ValueError(
            "No text found in this CV. If it's a scanned PDF, export it as a text-based PDF or DOCX instead."
        )
    return text[:MAX_CHARS]


def _pdf_text(data: bytes) -> str:
    with pdfplumber.open(io.BytesIO(data)) as pdf:
        return "\n".join(page.extract_text() or "" for page in pdf.pages)


def _docx_text(data: bytes) -> str:
    doc = Document(io.BytesIO(data))
    parts = [p.text for p in doc.paragraphs]
    # Many CV templates put skills / experience inside tables.
    for table in doc.tables:
        for row in table.rows:
            parts.extend(cell.text for cell in row.cells)
    return "\n".join(parts)


def _txt_text(data: bytes) -> str:
    try:
        return data.decode("utf-8-sig")
    except UnicodeDecodeError:
        return data.decode("latin-1")
